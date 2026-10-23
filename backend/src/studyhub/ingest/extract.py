"""Estrazione del testo dai documenti. I PDF NON vengono renderizzati in anticipo:
le immagini delle pagine si generano solo su richiesta (render_page) e finiscono in una
cache LRU limitata fuori dal Desktop."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

from ..config import LOCAL_DIR

PAGE_CACHE = LOCAL_DIR / "cache" / "pages"
PAGE_CACHE_MAX_BYTES = 500 * 1024 * 1024

CODE_EXT = {".c", ".h", ".cpp", ".hpp", ".java", ".py", ".m", ".v", ".sv", ".s", ".asm", ".sh", ".js", ".ts", ".r"}
TEXT_EXT = {".md", ".txt", ".csv"}
OFFICE_EXT = {".docx", ".pptx"}
SUPPORTED = {".pdf", *CODE_EXT, *TEXT_EXT, *OFFICE_EXT}

_MATH_CHARS = set("∑∫∮∂√∞≤≥≠≈≡∈∉⊂⊆∪∩∀∃∇±×÷→⇒⇔αβγδεζηθλμνξπρστφχψωΓΔΘΛΞΠΣΦΨΩ")
_MATH_FONT = re.compile(r"CMMI|CMSY|CMEX|MSBM|Math|Symbol|STIX|Cambria Math", re.I)


@dataclass
class Page:
    number: int          # 1-based
    text: str
    is_scan: bool = False
    is_visual: bool = False


@dataclass
class Extracted:
    kind: str
    title: str
    pages: list[Page] = field(default_factory=list)


def _pdf(path: Path) -> Extracted:
    doc = pymupdf.open(path)
    title = (doc.metadata or {}).get("title") or path.stem
    pages = []
    for p in doc:
        text = p.get_text("text").strip()
        chars = len(text)
        n_images = len(p.get_images(full=False))
        math_glyphs = sum(ch in _MATH_CHARS for ch in text)
        math_font = any(_MATH_FONT.search(f[3] or "") for f in p.get_fonts(full=False))
        try:
            n_drawings = len(p.get_drawings()) if chars < 4000 else 0
        except Exception:
            n_drawings = 0
        is_scan = chars < 40 and n_images > 0
        # "Visiva" = da mandare a Claude come immagine perché il testo estratto non basta:
        # scansioni, formule dense, schemi/figure con poco testo. Le pagine di solo testo no.
        equations = text.count("=") + math_glyphs
        is_visual = (is_scan
                     or (math_font and equations >= 6)
                     or math_glyphs >= 12
                     or n_drawings > 150
                     or (n_images > 0 and chars < 300))
        pages.append(Page(p.number + 1, text, is_scan, is_visual))
    doc.close()
    return Extracted("pdf", title.strip() or path.stem, pages)


def _docx(path: Path) -> Extracted:
    import docx
    d = docx.Document(str(path))
    parts = []
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(("#" * int(para.style.name[-1]) + " " + t) if para.style.name.startswith("Heading") and para.style.name[-1].isdigit() else t)
    return Extracted("docx", path.stem, [Page(1, "\n".join(parts))])


def _pptx(path: Path) -> Extracted:
    from pptx import Presentation
    prs = Presentation(str(path))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        pages.append(Page(i, "\n".join(texts), is_visual=any(sh.shape_type == 13 for sh in slide.shapes)))
    return Extracted("pptx", path.stem, pages)


def extract(path: Path) -> Extracted:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    text = path.read_text(encoding="utf-8", errors="replace")
    kind = "code" if ext in CODE_EXT else ("md" if ext == ".md" else "text")
    return Extracted(kind, path.name, [Page(1, text)])


def render_page(pdf: Path, page: int, key: str, dpi: int = 120) -> Path:
    """PNG di una pagina, generato solo quando serve (UI, OCR, Claude) e tenuto in cache."""
    PAGE_CACHE.mkdir(parents=True, exist_ok=True)
    out = PAGE_CACHE / f"{key[:20]}_{page}_{dpi}.png"
    if out.exists():
        out.touch()  # aggiorna l'età per la LRU
        return out
    doc = pymupdf.open(pdf)
    try:
        pix = doc[page - 1].get_pixmap(dpi=dpi)
        tmp = out.with_suffix(".part")
        pix.save(str(tmp), output="png")
        tmp.replace(out)
    finally:
        doc.close()
    _prune_cache()
    return out


def _prune_cache() -> None:
    files = sorted(PAGE_CACHE.glob("*.png"), key=lambda f: f.stat().st_mtime)
    total = sum(f.stat().st_size for f in files)
    while files and total > PAGE_CACHE_MAX_BYTES:
        f = files.pop(0)
        total -= f.stat().st_size
        f.unlink(missing_ok=True)
